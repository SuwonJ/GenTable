from __future__ import annotations

import re
import io
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.ingestion import build_time_mask, extract_credits, extract_grades, parse_meetings, to_minutes


DAY_ORDER = ["월", "화", "수", "목", "금"]
DAY_MAP = {day: index for index, day in enumerate(DAY_ORDER)}


def load_catalog(table_dir: Path) -> dict:
    """내부 tables 폴더를 읽기 위한 기존 호환 로더."""
    files = list(table_dir.glob("*.xlsx")) + list(table_dir.glob("*.csv"))
    return process_files(files)


def process_files(files: list[Path | io.BytesIO], filenames: list[str] | None = None) -> dict:
    courses = []
    curriculum_required = {"1": {"1": [], "2": []}, "2": {"1": [], "2": []}, "3": {"1": [], "2": []}, "4": {"1": [], "2": []}}
    warnings = []

    if filenames is None:
        filenames = [f.name if hasattr(f, "name") else f"file_{i}" for i, f in enumerate(files)]

    for file, name in zip(files, filenames):
        # 교과과정 파일은 강의 목록이 아니라 학년별 필수 과목으로 분리해 읽는다.
        if "교과과정" in name:
            try:
                curriculum_required = load_curriculum_required(file)
            except Exception as exc:
                warnings.append(f"{name}: 교과과정 로드 실패 ({exc})")
            continue

        try:
            if name.endswith(".csv"):
                df = pd.read_csv(file)
            else:
                try:
                    # 학교 엑셀 원본은 데이터 시트가 있는 경우가 많아 먼저 시도한다.
                    df = pd.read_excel(file, sheet_name="Data")
                except Exception:
                    df = pd.read_excel(file)
            
            normalized, stats = normalize_courses(df, source=name)
            if not normalized:
                warnings.append(f"{name}: 유효한 강의 데이터가 없거나 컬럼 형식이 맞지 않습니다.")
            skipped_count = stats["missing_name"] + stats["invalid_time"]
            if skipped_count:
                warnings.append(
                    f"{name}: {skipped_count}개 행이 제외되었습니다 "
                    f"(과목명 없음 {stats['missing_name']}개, 시간표 파싱 실패 {stats['invalid_time']}개)."
                )
            courses.extend(normalized)
        except Exception as exc:
            warnings.append(f"{name}: 읽기 실패 ({exc})")
            continue

    return {
        "courses": courses,
        "curriculum_required": curriculum_required,
        "warnings": warnings,
    }


def normalize_courses(df: pd.DataFrame, source: str) -> tuple[list[dict], dict]:
    # 원본 엑셀 컬럼명을 내부에서 쓰는 키 이름으로 바꾼다.
    mapping = {
        "이수구분(주전공)": "category_raw",
        "과목번호": "course_id",
        "과목명": "name",
        "교수명": "professor",
        "시간/학점(설계)": "credit_raw",
        "여석": "seats",
        "강의시간(강의실)": "time_raw",
        "수강대상": "target_raw",
        "교과영역": "area",
    }
    
    # 이미 정규화된 쉼표 구분 파일도 같은 로직으로 처리할 수 있게 매핑한다.
    col_map = {}
    for target, source_key in mapping.items():
        if target in df.columns:
            col_map[target] = source_key
        elif source_key in df.columns:
            col_map[source_key] = source_key

    renamed = df.rename(columns=col_map)

    required_cols = ["course_id", "name", "category_raw", "time_raw"]
    for col in required_cols:
        if col not in renamed.columns:
            # 과목번호 컬럼명이 깨진 경우 값 패턴으로 한 번 더 찾는다.
            found = False
            for c in renamed.columns:
                if col == "course_id" and renamed[c].astype(str).str.match(r"^\d{6}-\d{2}$").any():
                    renamed = renamed.rename(columns={c: "course_id"})
                    found = True
                    break
            if not found:
                return [], {"missing_name": 0, "invalid_time": 0}

    courses = []
    stats = {"missing_name": 0, "invalid_time": 0}
    for row in renamed.fillna("").to_dict("records"):
        name = str(row["name"]).strip()
        if not name:
            stats["missing_name"] += 1
            continue

        category_raw = str(row.get("category_raw", "")).strip()
        # 시간 정보가 없는 행은 시간표 생성에 사용할 수 없으므로 제외한다.
        meetings = parse_meetings(str(row.get("time_raw", "")).strip())
        if not meetings:
            stats["invalid_time"] += 1
            continue

        credits = extract_credits(str(row.get("credit_raw", "")))
        seats = safe_int(row.get("seats", 0))
        course = {
            "source": source,
            "course_id": str(row.get("course_id", "")).strip(),
            "name": name,
            "category_raw": category_raw,
            "category": categorize(category_raw),
            "category_label": category_label(category_raw),
            "area": str(row.get("area", "")).strip(),
            "professor": str(row.get("professor", "")).strip(),
            "credits": credits,
            "seats": seats,
            "time_raw": str(row.get("time_raw", "")).strip(),
            "target_raw": str(row.get("target_raw", "")).strip(),
            "target_grades": extract_grades(str(row.get("target_raw", "")).strip()),
            "meetings": meetings,
        }
        # 분 단위 비트마스크를 미리 만들어 시간 충돌 검사를 빠르게 한다.
        course["time_mask"] = build_time_mask(meetings)
        courses.append(course)

    return courses, stats


def load_curriculum_required(file: Path | io.BytesIO) -> dict:
    result = {"1": {"1": [], "2": []}, "2": {"1": [], "2": []}, "3": {"1": [], "2": []}, "4": {"1": [], "2": []}}
    
    # 업로드 파일은 읽기 위치를 처음으로 되돌린 뒤 엑셀 라이브러리에 넘긴다.
    if isinstance(file, io.BytesIO):
        file.seek(0)
        
    wb = load_workbook(file, data_only=True)
    ws = wb.active

    for row in range(6, ws.max_row + 1):
        grade = ws.cell(row=row, column=1).value
        if grade is None:
            continue
        grade_text = str(grade).strip()
        if not grade_text.isdigit():
            continue
        for semester, offset in [("1", 2), ("2", 13)]:
            category = str(ws.cell(row=row, column=offset).value or "").strip()
            name = str(ws.cell(row=row, column=offset + 2).value or "").strip()
            # 생성에 자동 반영할 필수 과목만 교과과정에서 추출한다.
            if not name or category not in {"전필", "교필"}:
                continue
            if grade_text in result:
                result[grade_text][semester].append({"name": name, "category": category})
    return result


def categorize(raw: str) -> str:
    raw = str(raw)
    if "전필" in raw:
        return "major_required"
    if "교필" in raw:
        return "general_required"
    if "교선" in raw:
        return "general_elective"
    if "전선" in raw:
        return "major_elective"
    return "other"


def category_label(raw: str) -> str:
    if raw:
        return str(raw)
    return "미분류"


def safe_int(value: object) -> int:
    try:
        return int(float(value))
    except Exception:
        return 0
